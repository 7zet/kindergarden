from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.billing.models import Invoice, Payment

from .analytics import attendance_daily, financial_snapshot, group_profitability
from .models import Expense, PayrollEntry


def _style_sheet(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F5F")
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(45, max(12, max(len(str(c.value or "")) for c in column) + 2))


def management_excel(kg, period):
    snap = financial_snapshot(kg, period)
    wb = Workbook()
    ws = wb.active; ws.title = "Umumiy"
    ws.append(["Ko'rsatkich", "Qiymat"])
    for label, value in [("Davr", period.strftime("%m.%Y")), ("Reja", snap["charged"]),
                         ("Tushum", snap["income"]), ("Xarajat", snap["expenses"]),
                         ("Ish haqi", snap["payroll"]), ("Sof natija", snap["profit"]),
                         ("Reja bajarilishi %", snap["plan_rate"])]: ws.append([label, value])
    _style_sheet(ws)
    sheets = [
        ("To'lovlar", ["Sana", "Bola", "Usul", "Summa", "Izoh"],
         ([timezone.localtime(x.received_at).strftime("%d.%m.%Y %H:%M"),
           x.child.full_name if x.child else "", x.get_method_display(), x.amount, x.note]
          for x in Payment.objects.filter(kindergarten=kg, is_reversed=False,
              received_at__date__range=(snap["start"], snap["end"])).select_related("child"))),
        ("Xarajatlar", ["Sana", "Kategoriya", "Nomi", "Yetkazib beruvchi", "Usul", "Summa"],
         ([timezone.localtime(x.paid_at).strftime("%d.%m.%Y %H:%M"), x.category.name,
           x.title, x.supplier.name if x.supplier else "", x.get_method_display(), x.amount]
          for x in Expense.objects.filter(kindergarten=kg, is_void=False,
              paid_at__date__range=(snap["start"], snap["end"])).select_related("category", "supplier"))),
        ("Ish haqi", ["Xodim", "Asosiy", "Avans", "Bonus", "Jarima", "Sof", "Holat"],
         ([x.employee.full_name, x.base_salary, x.advance, x.bonus, x.penalty,
           x.net_amount, x.get_status_display()] for x in PayrollEntry.objects.filter(
               kindergarten=kg, period=snap["period"]).select_related("employee"))),
        ("Qarzlar", ["Bola", "Invoys", "Davr", "Muddat", "Qoldiq"],
         ([x.enrollment.child.full_name, x.number, x.period.strftime("%m.%Y"),
           x.due_at.strftime("%d.%m.%Y"), x.balance] for x in Invoice.objects.filter(
               kindergarten=kg).unpaid().select_related("enrollment__child"))),
        ("Guruhlar", ["Guruh", "Reja", "Tushum", "Bevosita xarajat", "Marja"],
         ([x["group"].name, x["charged"], x["collected"], x["direct_cost"], x["margin"]]
          for x in group_profitability(kg, period))),
        ("Davomat", ["Sana", "Keldi", "Kelmadi", "Check-in", "Check-out"],
         ([x["day"].strftime("%d.%m.%Y"), x["present"], x["absent"],
           x["checkins"], x["checkouts"]] for x in attendance_daily(kg, period))),
    ]
    for title, headers, rows in sheets:
        sheet = wb.create_sheet(title); sheet.append(headers)
        for row in rows: sheet.append(row)
        _style_sheet(sheet)
    output = BytesIO(); wb.save(output); return output.getvalue()


def management_pdf(kg, period):
    snap = financial_snapshot(kg, period)
    output = BytesIO()
    font = "Helvetica"
    for path in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            pdfmetrics.registerFont(TTFont("ReportSans", path)); font = "ReportSans"; break
        except Exception: pass
    styles = getSampleStyleSheet()
    for style in styles.byName.values(): style.fontName = font
    story = [Paragraph(f"{kg.legal_name or kg.name}", styles["Heading2"]),
             Paragraph(f"BOSHQARUV HISOBOTI · {period:%m.%Y}", styles["Title"]), Spacer(1, 6 * mm)]
    rows = [["Ko'rsatkich", "Summa / qiymat"],
            ["Hisoblangan reja", f"{snap['charged']:,.0f} so'm"],
            ["Amaldagi tushum", f"{snap['income']:,.0f} so'm"],
            ["Xarajatlar", f"{snap['expenses']:,.0f} so'm"],
            ["Ish haqi", f"{snap['payroll']:,.0f} so'm"],
            ["Sof operatsion natija", f"{snap['profit']:,.0f} so'm"],
            ["Reja bajarilishi", f"{snap['plan_rate']:.1f}%"],
            ["Davomat belgilari", str(snap["present"])],
            ["Check-in / check-out", f"{snap['checkins']} / {snap['checkouts']}"]]
    table = Table(rows, colWidths=[95 * mm, 70 * mm])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173f5f")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTNAME", (0, 0), (-1, -1), font),
                               ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#ccd5e0")),
                               ("PADDING", (0, 0), (-1, -1), 7)]))
    story.append(table); story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("Guruhlar rentabelligi", styles["Heading2"]))
    grows = [["Guruh", "Tushum", "Xarajat", "Marja"]] + [
        [r["group"].name, f"{r['collected']:,.0f}", f"{r['direct_cost']:,.0f}", f"{r['margin']:,.0f}"]
        for r in group_profitability(kg, period)]
    gt = Table(grows, colWidths=[65 * mm, 35 * mm, 35 * mm, 35 * mm])
    gt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173f5f")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, -1), font),
                            ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#ccd5e0")),
                            ("PADDING", (0, 0), (-1, -1), 6)]))
    story.append(gt)
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    doc.build(story); return output.getvalue()
